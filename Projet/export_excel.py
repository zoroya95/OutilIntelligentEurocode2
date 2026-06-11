from openpyxl import Workbook


def exporter_resultats(resultat, nom_feuille, fichier):

    wb = Workbook()
    ws = wb.active
    ws.title = nom_feuille

    ligne = 1

    for cle, valeur in resultat.items():
        ws.cell(row=ligne, column=1, value=cle)
        ws.cell(row=ligne, column=2, value=str(valeur))
        ligne += 1

    wb.save(fichier)